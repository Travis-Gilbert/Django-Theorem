"""Replace docx-01 with a real merged-table DOCX and LibreOffice-rendered PDF.

Run after generate_corpus.py. LibreOffice is a fixture authoring tool, not an
ingestion dependency. No PDF/HTML approximation substitutes for its renderer.
"""
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory


def render(corpus: Path, soffice: Path):
    from docx import Document
    from docx.shared import Inches, Pt
    import pypdfium2 as pdfium

    manifest_path = corpus / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    row = next(r for r in manifest if r["id"] == "docx-01")
    gold_path = corpus / row["gold"]
    gold = json.loads(gold_path.read_text())
    source = corpus / row["source"]
    # Derive the authored content from existing gold fields, not parser output.
    fields = {key: value["value"] for key, value in gold["fields"].items()}
    title = f"Lease record {fields['reference']}"
    lines = [f"reference: {fields['reference']}", f"person: {fields['person']}",
             f"amount: {fields['amount']}", f"deadline: {fields['deadline']}",
             f"{fields['person']} approved the lease."]
    document = Document()
    document.core_properties.created = document.core_properties.modified = datetime(2026, 9, 5)
    document.core_properties.title = title
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    document.add_heading(title, level=1)
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Payment schedule"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Rent"
    table.cell(1, 1).text, table.cell(1, 2).text = str(fields["amount"]), "monthly"
    table.cell(2, 1).text, table.cell(2, 2).text = str(fields["amount"] * 12), "annual"
    document.save(source)
    version = subprocess.run([str(soffice), "--version"], check=True, capture_output=True, text=True, timeout=30).stdout.strip()
    with TemporaryDirectory(prefix="theorem-libreoffice-") as directory:
        temporary = Path(directory)
        command = [str(soffice), "--headless", "--norestore", "--nologo",
                   "-env:UserInstallation=" + (temporary / "profile").as_uri(),
                   "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(temporary), str(source.resolve())]
        conversion = subprocess.run(command, check=True, capture_output=True, text=True, timeout=120)
        rendered = temporary / (source.stem + ".pdf")
        if not rendered.exists():
            raise RuntimeError("LibreOffice produced no PDF: " + conversion.stderr)
        pdf_path = corpus / row["comparison_pdf"]
        pdf_path.write_bytes(rendered.read_bytes())
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        if len(pdf) != 1:
            raise ValueError("The paired rendered fixture must be one page")
        page = pdf[0]
        try:
            textpage = page.get_textpage()
            try:
                rendered_text = textpage.get_text_range()
                if any(text not in rendered_text for text in ["Payment schedule", "Rent", "monthly", "annual"]):
                    raise ValueError("The real rendered PDF lost authored table contents")
            finally:
                textpage.close()
            bitmap = page.render(scale=1.5)
            try:
                bitmap.to_pil().save(corpus / row["page_image"])
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        pdf.close()
    gold.update(texts=[title, *lines, "Payment schedule", "Rent", str(fields["amount"]),
                       "monthly", str(fields["amount"] * 12), "annual"],
        tables=[[["Payment schedule", "", ""], ["Rent", str(fields["amount"]), "monthly"],
                 ["", str(fields["amount"] * 12), "annual"]]],
        merged_cells=[[{"row": 0, "column": 0, "row_span": 1, "col_span": 3},
                       {"row": 1, "column": 0, "row_span": 2, "col_span": 1}]])
    gold_path.write_text(json.dumps(gold, indent=2) + "\n")
    for key in ("source", "comparison_pdf", "page_image", "gold"):
        row[key + "_sha256"] = sha256((corpus / row[key]).read_bytes()).hexdigest()
    receipt = {"renderer": "LibreOffice", "version": version, "filter": "writer_pdf_Export",
               "source": row["source"], "source_sha256": row["source_sha256"],
               "comparison_pdf": row["comparison_pdf"], "comparison_pdf_sha256": row["comparison_pdf_sha256"],
               "page_image": row["page_image"], "page_image_sha256": row["page_image_sha256"],
               "page_count": 1, "image_renderer": "pypdfium2", "image_scale": 1.5}
    receipt_path = corpus / "docx-01.render.json"
    receipt_path.write_text(json.dumps(receipt, indent=2) + "\n")
    row.update(page_image_provenance="libreoffice_docx_pdfium_rendered",
               renderer_receipt=receipt_path.name, renderer_receipt_sha256=sha256(receipt_path.read_bytes()).hexdigest())
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    return receipt


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--corpus", type=Path, default=Path(__file__).parent / "corpus")
    parser.add_argument("--soffice", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(render(args.corpus, args.soffice), indent=2))
